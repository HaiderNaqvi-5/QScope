"""Scan session endpoints."""
from datetime import datetime
import hashlib
from html import escape
import json
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse, Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db_session
from app.models import (Baseline, BehaviorCoverageRecord, Finding, GeneratedTestCase, ManualTestCase, Project,
                        RemediationAttempt, RemediationFixPack, Report, ScanEventRecord, ScanSession, ScanTaskRecord)
from app.schemas.projects import ProjectModel
from app.schemas.scans import PersistedScanEvent, ReportHistoryResponse, ScanSessionResponse, ScanStartRequest, ScanTaskResponse
from app.services.preflight import build_scan_plan
from app.services.architecture import build_architecture_map
from app.services.runtime import cancel_scan, event_stream, start_scan
from app.services.findings import score_findings
from app.services.release_gate import evaluate_release_gate
from app.schemas.findings import FindingResponse, ScanReportResponse

router = APIRouter()


_REPORT_SECTIONS = (
    ("executive_summary", "Executive Summary"), ("project_information", "Project Information"),
    ("architecture", "Detected Architecture & Technology Stack"),
    ("methodology", "Scope and Test Methodology"), ("tool_coverage", "Scan Completeness / Tool Coverage"),
    ("build", "Build & Compilation Analysis"), ("dependencies", "Dependency / SBOM Analysis"),
    ("code_quality", "Code Quality & Maintainability"), ("hygiene", "Code Hygiene & AI-Slop Review"),
    ("security", "Security Analysis — SAST / Secrets / Dependencies / DAST"), ("api", "API Testing"),
    ("functional", "Automated Functional Testing"), ("manual", "Manual Testing"),
    ("regression", "Regression Testing"), ("responsive", "Responsive Design Testing"),
    ("accessibility", "Accessibility Testing"), ("performance", "Lighthouse / Frontend Performance"),
    ("load", "Load & Latency Testing"), ("deprecation", "Deprecation Analysis"),
    ("test_health", "Test Coverage and Existing Test Health"), ("severity", "Findings by Severity"),
    ("remediation", "Prioritized Remediation Plan"), ("llm", "LLM-Assisted Review Summary"),
    ("gates", "Final Scores / Quality Gates"), ("conclusion", "Conclusion"),
    ("tool_versions", "Appendix — Tool Versions"), ("evidence", "Appendix — Evidence / Screenshots"),
    ("sbom", "Appendix — Dependency/SBOM Summary"),
    ("logical", "Logical / Business-Rule Testing"), ("authorization", "Role & Authorization Matrix"),
    ("edge", "Edge-Case / Boundary Testing"), ("contract", "Contract Testing"),
    ("data_integrity", "Data Integrity"), ("behavior", "Behavior Coverage / Test Gaps"),
    ("mutation", "Mutation / Test Strength"), ("concurrency", "Concurrency / Resilience"),
    ("fix_verify", "Fix / Verify Status"), ("timeline", "Evidence Timeline"),
)


def _report_sections(scan: ScanSession, project: Project | None, findings: list[Finding],
                     generated: list[GeneratedTestCase], manual: list[ManualTestCase],
                     remediations: list[RemediationAttempt], fix_packs: list[RemediationFixPack],
                     behavior: list[BehaviorCoverageRecord],
                     events: list[ScanEventRecord],
                     gate: dict, score: int) -> list[dict]:
    results = scan.results or []
    tokens = {
        "build": ("BUILD", "COMPILE", "TYPE"), "dependencies": ("DEPEND", "SBOM"),
        "code_quality": ("STATIC", "QUALITY", "LINT"), "hygiene": ("HYGIENE", "SLOP"),
        "security": ("SECURITY", "SAST", "SECRET", "DAST", "ZAP"), "api": ("API", "SCHEMA", "POSTMAN"),
        "functional": ("TEST", "FUNCTIONAL"), "responsive": ("RESPONSIVE", "VISUAL"),
        "accessibility": ("ACCESSIBILITY", "AXE"), "performance": ("PERFORMANCE", "LIGHTHOUSE"),
        "load": ("LOAD", "LATENCY", "K6", "JMETER"), "deprecation": ("DEPRECATION",),
        "test_health": ("TEST", "COVERAGE"), "llm": ("LLM", "AI_REVIEW"),
    }

    def task_evidence(key: str) -> list[dict]:
        wanted = tokens.get(key, ())
        return [item for item in results if any(token in str(item.get("stage", item.get("task_id", ""))).upper()
                                                for token in wanted)]

    generated_by = {category: [item for item in generated if item.category == category]
                    for category in {"LOGICAL", "EDGE", "CONTRACT", "DATA_INTEGRITY", "MUTATION",
                                     "CONCURRENCY", "RESILIENCE"}}
    sections: list[dict] = []
    for key, title in _REPORT_SECTIONS:
        evidence: list | dict = task_evidence(key)
        status = "NOT_APPLICABLE"
        summary = "No applicable persisted evidence was recorded."
        if evidence:
            statuses = {str(item.get("status", "UNKNOWN")) for item in evidence}
            status = "TOOL_MISSING" if "TOOL_MISSING" in statuses else (
                "TOOL_ERROR" if statuses & {"FAILED", "TOOL_ERROR", "TIMED_OUT"} else (
                    "NOT_APPLICABLE" if statuses == {"NOT_APPLICABLE"} else
                    "SKIPPED" if statuses <= {"SKIPPED", "NOT_APPLICABLE"} else "COMPLETED"))
            summary = f"{len(evidence)} persisted task result(s): {', '.join(sorted(statuses))}."
        if key in {"executive_summary", "gates", "conclusion"}:
            status, evidence = "COMPLETED", {"score": score, "decision": gate["decision"],
                                               "blockers": gate["blockers"]}
            summary = f"Quality score {score}/100; release decision {gate['decision']}."
        elif key in {"project_information", "architecture", "methodology"}:
            status = "COMPLETED" if project else "NOT_APPLICABLE"
            evidence = (project.project_model or {}) if project else {}
            summary = "Project discovery and scan-mode evidence." if project else summary
        elif key in {"tool_coverage", "tool_versions"}:
            status, evidence = "COMPLETED", results
            summary = f"{len(results)} task result(s) disclose execution and applicability status."
        elif key == "manual":
            evidence = [{"module": item.module, "status": item.status} for item in manual]
            status, summary = ("COMPLETED", f"{len(manual)} manual case(s) recorded.") if manual else (status, summary)
        elif key == "regression":
            status, evidence, summary = "COMPLETED", {}, "Baseline classifications are included with findings."
        elif key == "severity":
            counts = {level: sum(item.severity == level for item in findings)
                      for level in ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO")}
            status, evidence, summary = "COMPLETED", counts, f"{len(findings)} normalized finding(s)."
        elif key in {"remediation", "fix_verify"}:
            evidence = {"attempts": [{"status": item.status, "risk": item.risk, "outcome": item.final_outcome}
                                     for item in remediations],
                        "fix_packs": [{"root_cause_id": item.root_cause_id, "finding_ids": item.finding_ids,
                                       "confidence": item.correlation_confidence, "status": item.status}
                                      for item in fix_packs]}
            if remediations or fix_packs:
                status, summary = "COMPLETED", f"{len(remediations)} attempt(s), {len(fix_packs)} Fix Pack(s)."
        elif key in {"logical", "authorization", "edge", "contract", "data_integrity", "mutation"}:
            category = {"logical": "LOGICAL", "authorization": "LOGICAL", "edge": "EDGE",
                        "contract": "CONTRACT", "data_integrity": "DATA_INTEGRITY", "mutation": "MUTATION"}[key]
            rows = generated_by[category]
            evidence = [{"title": item.title, "actor": item.actor, "status": item.status,
                         "expected": item.expected, "evidence": item.evidence} for item in rows]
            status, summary = ("COMPLETED", f"{len(rows)} persisted scenario(s).") if rows else (status, summary)
        elif key == "behavior":
            evidence = [{"title": item.title, "criticality": item.criticality,
                         "coverage_status": item.coverage_status, "source": item.source} for item in behavior]
            status, summary = (("COMPLETED", f"{len(behavior)} behavior coverage record(s).")
                               if behavior else (status, summary))
        elif key == "concurrency":
            rows = generated_by["CONCURRENCY"] + generated_by["RESILIENCE"]
            evidence = [{"title": item.title, "status": item.status, "evidence": item.evidence} for item in rows]
            status, summary = ("COMPLETED", f"{len(rows)} bounded dynamic scenario(s).") if rows else (status, summary)
        elif key in {"evidence", "timeline"}:
            evidence = [{"event": item.event, "task_id": item.task_id,
                         "created_at": item.created_at.isoformat() if item.created_at else None} for item in events]
            status, summary = ("COMPLETED", f"{len(events)} durable event(s).") if events else (status, summary)
        elif key in {"sbom"} and project:
            evidence = {"dependency_count": len((project.project_model or {}).get("dependencies", []))}
            status, summary = "COMPLETED", "Dependency inventory summary included."
        sections.append({"key": key, "title": title, "status": status, "summary": summary,
                         "evidence": evidence})
    return sections


def _sections_html(sections: list[dict]) -> str:
    return "".join(
        f"<section><h2>{escape(section['title'])}</h2>"
        f"<p><strong>Status: {escape(section['status'])}</strong> — {escape(section['summary'])}</p>"
        f"<details><summary>Persisted evidence summary</summary><ul>"
        f"{''.join(f'<li>{escape(line)}</li>' for line in _evidence_digest(section['evidence']))}"
        f"</ul></details></section>"
        for section in sections
    )


def _evidence_digest(evidence: object, limit: int = 10) -> list[str]:
    """Return a bounded, readable report summary without embedding raw logs."""
    def compact(value: object) -> str:
        text = json.dumps(value, default=str, ensure_ascii=True) if isinstance(value, (dict, list)) else str(value)
        return text.replace("\n", " ")[:360]

    if isinstance(evidence, dict):
        lines = []
        for key, value in list(evidence.items())[:limit]:
            if isinstance(value, list):
                preview = ", ".join(
                    str(item.get("value") or item.get("name") or item.get("task_id") or item.get("title") or "record")
                    if isinstance(item, dict) else str(item)
                    for item in value[:4]
                )
                suffix = f"; e.g. {preview}" if preview else ""
                lines.append(f"{key}: {len(value)} item(s){suffix}")
            elif isinstance(value, dict):
                lines.append(f"{key}: {len(value)} field(s)")
            else:
                lines.append(f"{key}: {compact(value)}")
        return lines or ["No additional persisted evidence."]
    if isinstance(evidence, list):
        lines = []
        for item in evidence[:limit]:
            if isinstance(item, dict):
                identity = item.get("task_id") or item.get("title") or item.get("event") or "evidence"
                status = item.get("status")
                detail = item.get("tool") or item.get("stage") or item.get("summary")
                lines.append(" · ".join(str(part) for part in (identity, status, detail) if part))
            else:
                lines.append(compact(item))
        if len(evidence) > limit:
            lines.append(f"{len(evidence) - limit} additional item(s) retained in the scan record.")
        return lines or ["No additional persisted evidence."]
    return [compact(evidence)] if evidence else ["No additional persisted evidence."]


def _response(scan: ScanSession) -> ScanSessionResponse:
    approval_required = any(task.get("requires_user_confirmation") for task in (scan.plan or []))
    return ScanSessionResponse(
        id=scan.id, project_id=scan.project_id, mode=scan.mode, status=scan.status or "PENDING",
        approved=scan.approved == "true", approval_required=approval_required,
        results=scan.results or [], error=scan.error, created_at=scan.created_at,
        started_at=scan.started_at, completed_at=scan.completed_at,
        git_commit=scan.git_commit, git_branch=scan.git_branch,
        project_fingerprint=scan.project_fingerprint, overall_score=scan.overall_score,
        is_complete_audit=scan.is_complete_audit, release_readiness=scan.release_readiness,
        release_blockers=scan.release_blockers or [],
    )


@router.post("/projects/{project_id}/scans", response_model=ScanSessionResponse, status_code=202)
async def start_project_scan(project_id: str, request: ScanStartRequest, db: AsyncSession = Depends(get_db_session)) -> ScanSessionResponse:
    project = await db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    model = ProjectModel.model_validate(project.project_model).model_dump()
    if request.runtime_target:
        model["runtime_targets"] = [{
            "host": request.runtime_target.host,
            "port": request.runtime_target.port,
            "scheme": request.runtime_target.scheme,
            "base_url": request.runtime_target.base_url,
        }]
        project.project_model = model
    tools, tasks = build_scan_plan(project.id, model, request.mode)
    del tools
    plan = [task.model_dump() for task in tasks]
    approval_required = any(task.get("requires_user_confirmation") for task in plan)
    git = model.get("git", {})
    fingerprint = hashlib.sha256(json.dumps(model, sort_keys=True, default=str).encode()).hexdigest()
    session = ScanSession(id=str(uuid4()), project_id=project.id, mode=request.mode,
        status="AWAITING_APPROVAL" if approval_required else "PENDING",
        approved="false", plan=plan, results=[], git_commit=git.get("commit"), git_branch=git.get("branch"),
        project_fingerprint=fingerprint)
    db.add(session)
    db.add_all([ScanTaskRecord(id=str(uuid4()), scan_id=session.id, task_key=task["task_id"], stage=task["stage"],
        tool=task["tool"], adapter=task["adapter"], target=task.get("target", "."), status=task.get("status", "PENDING"),
        depends_on=task.get("depends_on", []), requires_confirmation=task.get("requires_user_confirmation", False)) for task in plan])
    await db.commit()
    await db.refresh(session)
    if not approval_required:
        await start_scan(session.id, project.root_path, session.plan)
    return _response(session)


@router.post("/scans/{scan_id}/approve", response_model=ScanSessionResponse, status_code=202)
async def approve_scan(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> ScanSessionResponse:
    scan = await db.get(ScanSession, scan_id)
    if not scan:
        raise HTTPException(status_code=404, detail="Scan session not found")
    if scan.status != "AWAITING_APPROVAL":
        raise HTTPException(status_code=409, detail="Scan does not require approval")
    project = await db.get(Project, scan.project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    scan.approved = "true"
    scan.status = "PENDING"
    await db.commit()
    await db.refresh(scan)
    await start_scan(scan.id, project.root_path, scan.plan, approved=True)
    return _response(scan)


@router.get("/scans/{scan_id}", response_model=ScanSessionResponse)
async def get_scan(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> ScanSessionResponse:
    scan = await db.get(ScanSession, scan_id)
    if not scan:
        raise HTTPException(status_code=404, detail="Scan session not found")
    return _response(scan)


@router.get("/scans/{scan_id}/tasks", response_model=list[ScanTaskResponse])
async def scan_tasks(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> list[ScanTaskResponse]:
    if not await db.get(ScanSession, scan_id):
        raise HTTPException(status_code=404, detail="Scan session not found")
    rows = (await db.execute(select(ScanTaskRecord).where(ScanTaskRecord.scan_id == scan_id))).scalars().all()
    return [ScanTaskResponse.model_validate(row) for row in rows]


@router.get("/scans/{scan_id}/event-history", response_model=list[PersistedScanEvent])
async def scan_event_history(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> list[PersistedScanEvent]:
    if not await db.get(ScanSession, scan_id):
        raise HTTPException(status_code=404, detail="Scan session not found")
    rows = (await db.execute(select(ScanEventRecord).where(ScanEventRecord.scan_id == scan_id)
                             .order_by(ScanEventRecord.created_at, ScanEventRecord.id))).scalars().all()
    return [PersistedScanEvent.model_validate(row) for row in rows]


@router.get("/projects/{project_id}/scans", response_model=list[ScanSessionResponse])
async def list_project_scans(project_id: str, db: AsyncSession = Depends(get_db_session)) -> list[ScanSessionResponse]:
    if not await db.get(Project, project_id):
        raise HTTPException(status_code=404, detail="Project not found")
    scans = (await db.execute(
        select(ScanSession).where(ScanSession.project_id == project_id).order_by(ScanSession.created_at.desc())
    )).scalars().all()
    return [_response(scan) for scan in scans]


@router.post("/scans/{scan_id}/cancel", response_model=ScanSessionResponse)
async def stop_scan(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> ScanSessionResponse:
    scan = await db.get(ScanSession, scan_id)
    if not scan:
        raise HTTPException(status_code=404, detail="Scan session not found")
    if scan.status in {"COMPLETED", "FAILED", "CANCELLED", "ERROR"}:
        return _response(scan)
    if scan.status == "AWAITING_APPROVAL":
        scan.status = "CANCELLED"
        await db.commit()
        await db.refresh(scan)
        return _response(scan)
    if not cancel_scan(scan_id):
        raise HTTPException(status_code=409, detail="Scan is not currently running")
    scan.status = "CANCELLING"
    await db.commit()
    await db.refresh(scan)
    return _response(scan)


@router.get("/scans/{scan_id}/events")
async def scan_events(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> StreamingResponse:
    if not await db.get(ScanSession, scan_id):
        raise HTTPException(status_code=404, detail="Scan session not found")
    return StreamingResponse(event_stream(scan_id), media_type="text/event-stream")


@router.get("/scans/{scan_id}/report", response_model=ScanReportResponse)
async def scan_report(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> ScanReportResponse:
    scan = await db.get(ScanSession, scan_id)
    if not scan:
        raise HTTPException(status_code=404, detail="Scan session not found")
    findings = (await db.execute(select(Finding).where(Finding.scan_id == scan_id))).scalars().all()
    project_id = scan.project_id
    project = await db.get(Project, project_id)
    dependency_count = len((project.project_model or {}).get("dependencies", [])) if project else 0
    api_spec_count = len((project.project_model or {}).get("api_specs", [])) if project else 0
    graphql_spec_count = len((project.project_model or {}).get("graphql_specs", [])) if project else 0
    postman_count = len((project.project_model or {}).get("postman_collections", [])) if project else 0
    runtime_target_configured = any(
        task.get("task_id") in {"runtime-api", "runtime-graphql", "runtime-postman"} and task.get("target") != "unconfigured"
        for task in (scan.plan or [])
    )
    api_task_present = any(task.get("task_id") in {"runtime-api", "runtime-graphql", "runtime-postman"} for task in (scan.plan or []))
    api_source_count = api_spec_count + graphql_spec_count + postman_count
    api_testing_status = "READY" if api_source_count and runtime_target_configured and api_task_present else (
        "API_SOURCE_FOUND_TARGET_REQUIRED" if api_source_count else "NOT_CONFIGURED"
    )
    security_task = next((task for task in (scan.plan or []) if task.get("task_id") == "runtime-zap"), None)
    load_task = next((task for task in (scan.plan or []) if task.get("task_id") in {
        "runtime-k6", "runtime-jmeter", "runtime-qsscope-load"}), None)
    security_status = "READY" if security_task and security_task.get("target") != "unconfigured" else (
        "TARGET_REQUIRED" if security_task else "NOT_CONFIGURED"
    )
    load_status = "READY" if load_task and load_task.get("target") != "unconfigured" else (
        "TARGET_REQUIRED" if load_task else "NOT_CONFIGURED"
    )
    accessibility_task = next((task for task in (scan.plan or []) if task.get("task_id") == "browser-accessibility"), None)
    performance_task = next((task for task in (scan.plan or []) if task.get("task_id") == "browser-performance"), None)
    accessibility_status = "READY" if accessibility_task and accessibility_task.get("target") != "unconfigured" else (
        "TARGET_REQUIRED" if accessibility_task else "NOT_CONFIGURED"
    )
    performance_status = "READY" if performance_task and performance_task.get("target") != "unconfigured" else (
        "TARGET_REQUIRED" if performance_task else "NOT_CONFIGURED"
    )
    baseline = (await db.execute(
        select(Baseline).where(Baseline.project_id == project_id).order_by(Baseline.created_at.desc())
    )).scalars().first()
    baseline_fingerprints = set(baseline.fingerprints or []) if baseline else set()
    manual_cases = list((await db.execute(select(ManualTestCase).where(
        ManualTestCase.project_id == project_id).order_by(ManualTestCase.created_at))).scalars().all())
    manual_payload = [{"id": case.id, "module": case.module, "feature": case.feature,
        "preconditions": case.preconditions or [], "steps": case.steps or [],
        "expected_result": case.expected_result, "actual_result": case.actual_result,
        "status": case.status, "severity": case.severity, "notes": case.notes,
        "source": case.source} for case in manual_cases]
    remediation_attempts = list((await db.execute(select(RemediationAttempt).where(
        RemediationAttempt.scan_id == scan_id).order_by(RemediationAttempt.created_at))).scalars().all())
    remediation_payload = [{
        "id": item.id, "finding_id": item.finding_id, "attempt_number": item.attempt_number,
        "mode": item.mode, "status": item.status, "risk": item.risk,
        "finding_confidence": float(item.finding_confidence), "fix_confidence": float(item.fix_confidence),
        "impact_radius": item.impact_radius or {}, "plan": item.plan or {},
        "changed_files": item.changed_files or [], "verification_plan": item.verification_plan or [],
        "evidence": item.evidence or {}, "final_outcome": item.final_outcome,
    } for item in remediation_attempts]
    fix_packs = list((await db.execute(select(RemediationFixPack).where(
        RemediationFixPack.scan_id == scan_id).order_by(RemediationFixPack.created_at,
                                                        RemediationFixPack.id))).scalars().all())
    fix_pack_payload = [{"id": item.id, "root_cause_id": item.root_cause_id,
                         "finding_ids": item.finding_ids or [],
                         "correlation_confidence": float(item.correlation_confidence),
                         "proposed_shared_correction": item.proposed_shared_correction,
                         "risk": item.risk, "impact_radius": item.impact_radius or {},
                         "verification_scenarios": item.verification_scenarios or {},
                         "status": item.status} for item in fix_packs]
    generated_cases = list((await db.execute(select(GeneratedTestCase).where(
        GeneratedTestCase.scan_id == scan_id).order_by(GeneratedTestCase.category,
                                                        GeneratedTestCase.endpoint))).scalars().all())
    generated_payload = [{"id": item.id, "category": item.category, "title": item.title,
        "method": item.method, "endpoint": item.endpoint, "actor": item.actor,
        "expected": item.expected or {}, "status": item.status, "evidence": item.evidence or {},
        "source_path": item.source_path} for item in generated_cases]
    behavior_records = list((await db.execute(select(BehaviorCoverageRecord).where(
        BehaviorCoverageRecord.project_id == project_id,
        (BehaviorCoverageRecord.scan_id == scan_id) | (BehaviorCoverageRecord.scan_id.is_(None)),
    ).order_by(BehaviorCoverageRecord.criticality, BehaviorCoverageRecord.category,
               BehaviorCoverageRecord.title))).scalars().all())
    behavior_payload = [{
        "id": item.id, "scan_id": item.scan_id, "capability_key": item.capability_key,
        "category": item.category, "title": item.title, "criticality": item.criticality,
        "coverage_status": item.coverage_status, "source": item.source,
        "evidence": item.evidence or {},
    } for item in behavior_records]
    events = list((await db.execute(select(ScanEventRecord).where(
        ScanEventRecord.scan_id == scan_id).order_by(ScanEventRecord.created_at,
                                                      ScanEventRecord.id))).scalars().all())
    current_fingerprints = {finding.fingerprint for finding in findings}
    response_findings = []
    for finding in findings:
        response = FindingResponse.model_validate(finding, from_attributes=True)
        response.baseline_state = "EXISTING" if finding.fingerprint in baseline_fingerprints else "NEW"
        response_findings.append(response)
    finding_payloads = [finding.model_dump() for finding in response_findings]
    score = score_findings(finding_payloads, scan.results or [], full_audit=scan.mode == "FULL")
    gate = evaluate_release_gate(finding_payloads, scan.results or [], full_audit=scan.mode == "FULL")
    sections = _report_sections(scan, project, findings, generated_cases, manual_cases,
                                remediation_attempts, fix_packs, behavior_records, events, gate, score)
    grade = ("Excellent" if score >= 90 else "Strong" if score >= 80 else
             "Acceptable with Improvements" if score >= 70 else "Risky" if score >= 60 else
             "Poor / Major Issues" if score >= 40 else "Failing")
    return ScanReportResponse(
        scan_id=scan_id, status=scan.status or "PENDING",
        approved=scan.approved == "true",
        approval_required=any(task.get("requires_user_confirmation") for task in (scan.plan or [])),
        score=score, grade=grade, release_readiness=gate["decision"],
        release_blockers=gate["blockers"], is_complete_audit=gate["is_complete_audit"],
        incomplete_reasons=gate["incomplete_reasons"],
        findings=response_findings,
        task_count=len(scan.results or []),
        failed_tasks=sum(1 for result in (scan.results or []) if result.get("status") in {"FAILED", "TIMED_OUT", "TOOL_MISSING"}),
        new_findings=sum(1 for finding in response_findings if finding.baseline_state == "NEW"),
        existing_findings=sum(1 for finding in response_findings if finding.baseline_state == "EXISTING"),
        resolved_findings=len(baseline_fingerprints - current_fingerprints),
        dependency_count=dependency_count,
        api_spec_count=api_spec_count,
        api_collection_count=postman_count,
        api_testing_status=api_testing_status,
        security_testing_status=security_status,
        load_testing_status=load_status,
        accessibility_status=accessibility_status,
        performance_status=performance_status,
        manual_test_count=len(manual_cases),
        manual_test_failures=sum(case.status == "FAIL" for case in manual_cases),
        manual_tests=manual_payload,
        remediation_count=len(remediation_attempts),
        remediation_verified=sum(item.final_outcome == "VERIFIED_FIXED" for item in remediation_attempts),
        remediation_rolled_back=sum((item.evidence or {}).get("rollback", {}).get("status") == "ROLLED_BACK"
                                    for item in remediation_attempts),
        remediation_manual_review=sum(item.status == "MANUAL_REVIEW_REQUIRED" for item in remediation_attempts),
        remediations=remediation_payload,
        fix_pack_count=len(fix_packs),
        fix_packs=fix_pack_payload,
        generated_test_count=len(generated_cases),
        generated_test_passed=sum(item.status == "PASSED" for item in generated_cases),
        generated_test_failed=sum(item.status == "FAILED" for item in generated_cases),
        generated_test_blocked=sum(item.status in {"BLOCKED", "NOT_APPLICABLE"} for item in generated_cases),
        generated_tests=generated_payload,
        behavior_coverage_total=len(behavior_records),
        behavior_coverage_covered=sum(item.coverage_status == "COVERED" for item in behavior_records),
        behavior_coverage_partial=sum(item.coverage_status == "PARTIALLY_COVERED" for item in behavior_records),
        behavior_coverage_untested=sum(item.coverage_status == "UNTESTED" for item in behavior_records),
        behavior_coverage_unknown=sum(item.coverage_status == "UNKNOWN" for item in behavior_records),
        behavior_coverage=behavior_payload,
        project_information={"id": project.id, "name": project.name, "root_path": project.root_path,
                             "git_commit": scan.git_commit, "git_branch": scan.git_branch,
                             "scan_mode": scan.mode} if project else {},
        architecture=build_architecture_map(project.project_model or {}) if project else {},
        tool_coverage=[{"task_id": item.get("task_id"), "stage": item.get("stage"),
                        "tool": item.get("tool"), "status": item.get("status"),
                        "duration_ms": item.get("duration_ms")} for item in (scan.results or [])],
        evidence_timeline=[{"id": item.id, "event": item.event, "task_id": item.task_id,
                            "payload": item.payload or {}, "created_at": item.created_at}
                           for item in events],
        report_sections=sections,
    )


@router.get("/scans/{scan_id}/report/export")
async def export_scan_report(
    scan_id: str,
    format: str = Query("json", pattern="^(json|markdown|html|docx|pdf)$"),
    db: AsyncSession = Depends(get_db_session),
) -> Response:
    scan = await db.get(ScanSession, scan_id)
    if not scan:
        raise HTTPException(status_code=404, detail="Scan session not found")
    if scan.status not in {"COMPLETED", "FAILED", "ERROR", "CANCELLED"}:
        raise HTTPException(status_code=409, detail="Report export is available only after the scan reaches a terminal state")
    report = await scan_report(scan_id, db)

    async def record_successful_export() -> None:
        db.add(Report(id=str(uuid4()), scan_id=scan_id, format=format))
        await db.commit()

    if format == "json":
        body = json.dumps(report.model_dump(), indent=2, default=str)
        await record_successful_export()
        return Response(body, media_type="application/json", headers={
            "Content-Disposition": f'attachment; filename="qsscope-{scan_id}.json"',
        })
    if format == "html":
        from html import escape
        finding_rows = "".join(
            f"<tr><td>{escape(finding.severity or '')}</td>"
            f"<td>{escape(finding.status)}</td><td>{escape(finding.title)}</td>"
            f"<td>{escape(finding.message)}</td></tr>"
            for finding in report.findings
        ) or '<tr><td colspan="4">No normalized findings.</td></tr>'
        manual_rows = "".join(f"<tr><td>{escape(case['module'])}</td><td>{escape(case['status'])}</td>"
                              f"<td>{escape(case['expected_result'])}</td><td>{escape(case['actual_result'] or '')}</td></tr>"
                              for case in report.manual_tests) or '<tr><td colspan="4">No manual tests recorded.</td></tr>'
        remediation_rows = "".join(
            f"<tr><td>{escape(item['mode'])}</td><td>{escape(item['risk'])}</td>"
            f"<td>{escape(item['status'])}</td><td>{escape(item.get('final_outcome') or '')}</td></tr>"
            for item in report.remediations
        ) or '<tr><td colspan="4">No remediation attempts recorded.</td></tr>'
        behavior_rows = "".join(
            f"<tr><td>{escape(item['category'])}</td><td>{escape(item['criticality'])}</td>"
            f"<td>{escape(item['coverage_status'])}</td><td>{escape(item['title'])}</td></tr>"
            for item in report.behavior_coverage
        ) or '<tr><td colspan="4">No behavior coverage records.</td></tr>'
        body = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>QSScope Scan Report</title>
<style>body{{font:16px system-ui,sans-serif;max-width:1100px;margin:2rem auto;padding:0 1rem;color:#172033}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #d5dbe5;padding:.6rem;text-align:left}}
th{{background:#eef2f7}}</style></head><body>
<h1>QSScope Scan Report</h1><p>Scan: <code>{escape(report.scan_id)}</code></p>
<p>Status: <strong>{escape(report.status)}</strong> · Score: <strong>{report.score}/100</strong></p>
<p>New: {report.new_findings} · Existing: {report.existing_findings} · Resolved: {report.resolved_findings}</p>
<h2>Findings</h2><table><thead><tr><th>Severity</th><th>Status</th><th>Title</th><th>Message</th></tr></thead>
<tbody>{finding_rows}</tbody></table><h2>Manual Testing</h2><table><thead><tr><th>Module</th><th>Status</th><th>Expected</th><th>Actual</th></tr></thead>
<tbody>{manual_rows}</tbody></table><h2>Automated Remediation &amp; Verification</h2>
<p>Attempted: {report.remediation_count} · Verified: {report.remediation_verified} · Rolled back: {report.remediation_rolled_back} · Manual review: {report.remediation_manual_review}</p>
<table><thead><tr><th>Mode</th><th>Risk</th><th>State</th><th>Outcome</th></tr></thead><tbody>{remediation_rows}</tbody></table>
<h2>Behavior Coverage &amp; Test Gaps</h2><p>Covered: {report.behavior_coverage_covered} · Partial: {report.behavior_coverage_partial} · Untested: {report.behavior_coverage_untested} · Unknown: {report.behavior_coverage_unknown}</p>
<table><thead><tr><th>Category</th><th>Criticality</th><th>Coverage</th><th>Behavior</th></tr></thead><tbody>{behavior_rows}</tbody></table>
<h2>Full Audit Section Ledger</h2>{_sections_html(report.report_sections)}</body></html>"""
        await record_successful_export()
        return Response(body, media_type="text/html", headers={
            "Content-Disposition": f'attachment; filename="qsscope-{scan_id}.html"',
        })
    if format == "docx":
        from io import BytesIO
        from docx import Document

        document = Document()
        document.add_heading("QSScope Scan Report", level=1)
        document.add_paragraph(f"Scan: {report.scan_id}")
        document.add_paragraph(f"Status: {report.status} | Score: {report.score}/100")
        document.add_paragraph(
            f"New: {report.new_findings} | Existing: {report.existing_findings} | "
            f"Resolved: {report.resolved_findings}"
        )
        document.add_heading("Findings", level=2)
        if report.findings:
            table = document.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for cell, heading in zip(table.rows[0].cells, ("Severity", "Status", "Title", "Message")):
                cell.text = heading
            for finding in report.findings:
                cells = table.add_row().cells
                cells[0].text = finding.severity or ""
                cells[1].text = finding.status
                cells[2].text = finding.title
                cells[3].text = finding.message
        else:
            document.add_paragraph("No normalized findings.")
        document.add_heading("Manual Testing", level=2)
        if report.manual_tests:
            manual_table = document.add_table(rows=1, cols=4)
            manual_table.style = "Table Grid"
            for cell, heading in zip(manual_table.rows[0].cells, ("Module", "Status", "Expected", "Actual")):
                cell.text = heading
            for case in report.manual_tests:
                cells = manual_table.add_row().cells
                cells[0].text = case["module"]
                cells[1].text = case["status"]
                cells[2].text = case["expected_result"]
                cells[3].text = case.get("actual_result") or ""
        else:
            document.add_paragraph("No manual tests recorded.")
        document.add_heading("Automated Remediation & Verification", level=2)
        document.add_paragraph(
            f"Attempted: {report.remediation_count} | Verified: {report.remediation_verified} | "
            f"Rolled back: {report.remediation_rolled_back} | Manual review: {report.remediation_manual_review}"
        )
        if report.remediations:
            remediation_table = document.add_table(rows=1, cols=4)
            remediation_table.style = "Table Grid"
            for cell, heading in zip(remediation_table.rows[0].cells, ("Mode", "Risk", "State", "Outcome")):
                cell.text = heading
            for item in report.remediations:
                cells = remediation_table.add_row().cells
                cells[0].text = item["mode"]
                cells[1].text = item["risk"]
                cells[2].text = item["status"]
                cells[3].text = item.get("final_outcome") or ""
        else:
            document.add_paragraph("No remediation attempts recorded.")
        document.add_heading("Behavior Coverage & Test Gaps", level=2)
        document.add_paragraph(
            f"Covered: {report.behavior_coverage_covered} | Partial: {report.behavior_coverage_partial} | "
            f"Untested: {report.behavior_coverage_untested} | Unknown: {report.behavior_coverage_unknown}"
        )
        if report.behavior_coverage:
            behavior_table = document.add_table(rows=1, cols=4)
            behavior_table.style = "Table Grid"
            for cell, heading in zip(behavior_table.rows[0].cells,
                                     ("Category", "Criticality", "Coverage", "Behavior")):
                cell.text = heading
            for item in report.behavior_coverage:
                cells = behavior_table.add_row().cells
                cells[0].text = item["category"]
                cells[1].text = item["criticality"]
                cells[2].text = item["coverage_status"]
                cells[3].text = item["title"]
        document.add_heading("Full Audit Section Ledger", level=1)
        for section in report.report_sections:
            document.add_heading(section["title"], level=2)
            document.add_paragraph(f"Status: {section['status']}")
            document.add_paragraph(section["summary"])
            for evidence_line in _evidence_digest(section["evidence"]):
                document.add_paragraph(evidence_line, style="List Bullet")
        buffer = BytesIO()
        document.save(buffer)
        await record_successful_export()
        return Response(buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={
            "Content-Disposition": f'attachment; filename="qsscope-{scan_id}.docx"',
        })
    if format == "pdf":
        import asyncio
        import sys
        import tempfile
        from html import escape
        from pathlib import Path

        finding_rows = "".join(
            f"<tr><td>{escape(finding.severity or '')}</td>"
            f"<td>{escape(finding.status)}</td><td>{escape(finding.title)}</td>"
            f"<td>{escape(finding.message)}</td></tr>"
            for finding in report.findings
        ) or "<tr><td colspan='4'>No normalized findings.</td></tr>"
        manual_rows = "".join(f"<tr><td>{escape(case['module'])}</td><td>{escape(case['status'])}</td>"
                              f"<td>{escape(case['expected_result'])}</td><td>{escape(case['actual_result'] or '')}</td></tr>"
                              for case in report.manual_tests) or "<tr><td colspan='4'>No manual tests recorded.</td></tr>"
        remediation_rows = "".join(
            f"<tr><td>{escape(item['mode'])}</td><td>{escape(item['risk'])}</td>"
            f"<td>{escape(item['status'])}</td><td>{escape(item.get('final_outcome') or '')}</td></tr>"
            for item in report.remediations
        ) or "<tr><td colspan='4'>No remediation attempts recorded.</td></tr>"
        behavior_rows = "".join(
            f"<tr><td>{escape(item['category'])}</td><td>{escape(item['criticality'])}</td>"
            f"<td>{escape(item['coverage_status'])}</td><td>{escape(item['title'])}</td></tr>"
            for item in report.behavior_coverage
        ) or "<tr><td colspan='4'>No behavior coverage records.</td></tr>"
        html_body = f"""<!doctype html><html><head><meta charset="utf-8">
<style>body{{font:12px Arial,sans-serif;margin:32px;color:#172033}}
h1{{font-size:24px}}table{{border-collapse:collapse;width:100%}}
th,td{{border:1px solid #ccd3df;padding:6px;text-align:left;vertical-align:top}}
th{{background:#eef2f7}}</style></head><body>
<h1>QSScope Scan Report</h1><p>Scan: <code>{escape(report.scan_id)}</code></p>
<p>Status: <strong>{escape(report.status)}</strong> · Score: <strong>{report.score}/100</strong></p>
<p>New: {report.new_findings} · Existing: {report.existing_findings} · Resolved: {report.resolved_findings}</p>
<h2>Findings</h2><table><thead><tr><th>Severity</th><th>Status</th><th>Title</th><th>Message</th></tr></thead>
<tbody>{finding_rows}</tbody></table><h2>Manual Testing</h2><table><thead><tr><th>Module</th><th>Status</th><th>Expected</th><th>Actual</th></tr></thead>
<tbody>{manual_rows}</tbody></table><h2>Automated Remediation &amp; Verification</h2>
<p>Attempted: {report.remediation_count} · Verified: {report.remediation_verified} · Rolled back: {report.remediation_rolled_back} · Manual review: {report.remediation_manual_review}</p>
<table><thead><tr><th>Mode</th><th>Risk</th><th>State</th><th>Outcome</th></tr></thead><tbody>{remediation_rows}</tbody></table>
<h2>Behavior Coverage &amp; Test Gaps</h2><p>Covered: {report.behavior_coverage_covered} · Partial: {report.behavior_coverage_partial} · Untested: {report.behavior_coverage_untested} · Unknown: {report.behavior_coverage_unknown}</p>
<table><thead><tr><th>Category</th><th>Criticality</th><th>Coverage</th><th>Behavior</th></tr></thead><tbody>{behavior_rows}</tbody></table>
<h2>Full Audit Section Ledger</h2>{_sections_html(report.report_sections)}</body></html>"""
        try:
            with tempfile.TemporaryDirectory(prefix="qsscope-pdf-") as temp_dir:
                source = Path(temp_dir) / "report.html"
                destination = Path(temp_dir) / "report.pdf"
                source.write_text(html_body, encoding="utf-8")
                process = await asyncio.create_subprocess_exec(
                    sys.executable, "-m", "app.services.pdf_renderer", str(source), str(destination),
                    stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
                )
                try:
                    _stdout, stderr = await asyncio.wait_for(process.communicate(), timeout=25)
                except asyncio.TimeoutError:
                    process.kill()
                    await process.wait()
                    raise
                if process.returncode != 0 or not destination.is_file():
                    detail = stderr.decode("utf-8", errors="replace")[-500:]
                    raise RuntimeError(detail or "renderer did not create a PDF")
                pdf = destination.read_bytes()
        except (OSError, RuntimeError, asyncio.TimeoutError) as exc:
            raise HTTPException(
                status_code=503,
                detail=f"PDF renderer unavailable; install Playwright Chromium browsers: {exc.__class__.__name__}",
            ) from exc
        await record_successful_export()
        return Response(pdf, media_type="application/pdf", headers={
            "Content-Disposition": f'attachment; filename="qsscope-{scan_id}.pdf"',
        })
    lines = [
        "# QSScope Scan Report", "",
        f"- Scan: `{report.scan_id}`",
        f"- Status: **{report.status}**",
        f"- Score: **{report.score}/100**",
        f"- New findings: **{report.new_findings}**",
        f"- Existing findings: **{report.existing_findings}**", "",
        f"- Resolved findings since baseline: **{report.resolved_findings}**", "",
        f"- Dependencies inventoried: **{report.dependency_count}**", "",
        f"- API specs: **{report.api_spec_count}**",
        f"- Postman collections: **{report.api_collection_count}**",
        f"- API testing readiness: **{report.api_testing_status}**", "",
        f"- Security testing readiness: **{report.security_testing_status}**",
        f"- Load testing readiness: **{report.load_testing_status}**", "",
        f"- Accessibility readiness: **{report.accessibility_status}**",
        f"- Performance readiness: **{report.performance_status}**", "",
        f"- Manual tests: **{report.manual_test_count}** ({report.manual_test_failures} failed)", "",
        "## Findings",
    ]
    if report.findings:
        lines.extend(
            f"- **{finding.severity} / {finding.status}** {finding.title}"
            f" — {(finding.file_path + ':' + (finding.line or '')) if finding.file_path else 'project'}"
            f" — {finding.message}"
            for finding in report.findings
        )
    else:
        lines.append("- No normalized findings.")
    lines.extend(["", "## Manual Testing"])
    if report.manual_tests:
        lines.extend(f"- **{case['status']}** {case['module']}: {case['expected_result']}"
                     for case in report.manual_tests)
    else:
        lines.append("- No manual tests recorded.")
    lines.extend(["", "## Automated Remediation & Verification",
                  f"- Attempted: **{report.remediation_count}**",
                  f"- Verified: **{report.remediation_verified}**",
                  f"- Rolled back: **{report.remediation_rolled_back}**",
                  f"- Manual review: **{report.remediation_manual_review}**"])
    if report.remediations:
        lines.extend(f"- **{item['status']}** {item['mode']} / {item['risk']} — {item.get('final_outcome') or 'pending'}"
                     for item in report.remediations)
    else:
        lines.append("- No remediation attempts recorded.")
    lines.extend(["", "## Behavior Coverage & Test Gaps",
                  f"- Covered: **{report.behavior_coverage_covered}**",
                  f"- Partially covered: **{report.behavior_coverage_partial}**",
                  f"- Untested: **{report.behavior_coverage_untested}**",
                  f"- Unknown: **{report.behavior_coverage_unknown}**"])
    if report.behavior_coverage:
        lines.extend(f"- **{item['coverage_status']} / {item['criticality']}** {item['category']}: {item['title']}"
                     for item in report.behavior_coverage)
    else:
        lines.append("- No behavior coverage records.")
    lines.extend(["", "# Full Audit Section Ledger"])
    for section in report.report_sections:
        lines.extend(["", f"## {section['title']}", "", f"**Status: {section['status']}**",
                      "", section["summary"]])
        lines.extend(f"- {line}" for line in _evidence_digest(section["evidence"]))
    await record_successful_export()
    return Response("\n".join(lines) + "\n", media_type="text/markdown", headers={
        "Content-Disposition": f'attachment; filename="qsscope-{scan_id}.md"',
    })


@router.get("/projects/{project_id}/reports", response_model=list[ReportHistoryResponse])
async def list_project_reports(
    project_id: str,
    db: AsyncSession = Depends(get_db_session),
) -> list[ReportHistoryResponse]:
    if not await db.get(Project, project_id):
        raise HTTPException(status_code=404, detail="Project not found")
    reports = (await db.execute(
        select(Report)
        .join(ScanSession, Report.scan_id == ScanSession.id)
        .where(ScanSession.project_id == project_id)
        .order_by(Report.created_at.desc())
    )).scalars().all()
    return [
        ReportHistoryResponse(id=report.id, scan_id=report.scan_id, format=report.format or "unknown", created_at=report.created_at)
        for report in reports
    ]


@router.post("/scans/{scan_id}/baseline", response_model=ScanReportResponse)
async def create_baseline(scan_id: str, db: AsyncSession = Depends(get_db_session)) -> ScanReportResponse:
    scan = await db.get(ScanSession, scan_id)
    if not scan:
        raise HTTPException(status_code=404, detail="Scan session not found")
    findings = (await db.execute(select(Finding).where(Finding.scan_id == scan_id))).scalars().all()
    db.add(Baseline(id=str(uuid4()), project_id=scan.project_id, fingerprints=[finding.fingerprint for finding in findings]))
    await db.commit()
    return await scan_report(scan_id, db)
